# views/dashboard.py

import json
from django.views import View
from django.views.generic import TemplateView
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.db.models import Sum
from django.db import transaction
from django.views.generic import ListView, DetailView, CreateView, UpdateView
from django.views.generic.edit import DeleteView
from django.urls import reverse_lazy
from django.contrib.auth import logout, authenticate, login
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib.auth.decorators import login_required
from django.contrib import messages
import openpyxl
from django.http import HttpResponse, JsonResponse
from django.db.models import Count, Q
import requests
from logistque.eglise_exports.views import export_eglises_to_excel, export_eglises_to_pdf, export_eglises_to_word
from logistque.forms import *
from logistque.models import *




def geocode_city(request):
    city = request.GET.get("q")
    if not city:
        return JsonResponse({"error": "City parameter missing"}, status=400)
    
    response = requests.get(
        "https://nominatim.openstreetmap.org/search",
        params={"q": city, "format": "json", "limit": 1},
        headers={"User-Agent": "VotreNomOuNomDuSite"}
    )
    
    if response.status_code != 200:
        return JsonResponse({"error": "Nominatim error"}, status=500)

    return JsonResponse(response.json(), safe=False)

# mixins.py
class ResponsableLogistiqueMixin(LoginRequiredMixin, UserPassesTestMixin):
    def test_func(self):
        return self.request.user.groups.filter(name='Responsable Logistique National').exists()
    

# eglise_views.py

class EgliseListView(LoginRequiredMixin, ListView):
    model = Eglise
    template_name = 'eglise/eglise_list.html'
    context_object_name = 'eglises'
    paginate_by = 10  # Nombre d'éléments par page

    def get_queryset(self):
        # Dans get_queryset()
        ville_id = self.request.GET.get('ville')
        
        queryset = super().get_queryset().select_related('ville', 'region', 'pasteur').order_by('-id')
        region_id = self.request.GET.get('region')
        
        if ville_id:
            queryset = queryset.filter(ville_id=ville_id).order_by('-id')
        
        if region_id:
            queryset = queryset.filter(
                Q(region_id=region_id) | Q(ville__region_id=region_id)).order_by('-id')
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['regions'] = Region.objects.all()
        context['selected_region'] = self.request.GET.get('region')
        return context

class EgliseDetailView(LoginRequiredMixin, DetailView):
    model = Eglise
    template_name = 'eglise/eglise_detail.html'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['regions'] = Region.objects.all()
        context['latitude'] = f"{self.object.latitude:.2f}" if self.object.latitude is not None else None
        context['longitude'] = f"{self.object.longitude:.2f}" if self.object.longitude is not None else None
        print('latitude', context['latitude'], 'et longitude', context['longitude'])
        return context

class EgliseCreateView(LoginRequiredMixin,CreateView):
    model = Eglise
    fields = ['nom', 'pasteur', 'phone', 'image','region','ville', 'pays', 'latitude', 'longitude']
    template_name = 'eglise/eglise_form.html'
    success_url = reverse_lazy('eglise-list')
    
    def form_valid(self, form):
        response = super().form_valid(form)

        # Option : extraire la géolocalisation transmise via JS
        latitude = self.request.POST.get('latitude')
        longitude = self.request.POST.get('longitude')
        if latitude and longitude:
            self.object.latitude = latitude
            self.object.longitude = longitude
            self.object.save()

        messages.success(self.request, f"✅ {self.object.nom}' enregistrée avec succès.")
        return response

    def form_invalid(self, form):
        messages.error(self.request, "❌ Veuillez corriger les erreurs dans le formulaire.")
        return super().form_invalid(form)
    
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
    
        # Ajout des données IP localisées dans le contexte pour affichage
        # context['ville'] = self.request.ville if hasattr(self.request, 'ville') else ''
        # context['pays'] = self.request.pays if hasattr(self.request, 'pays') else ''
        # context['latitude'] = self.request.latitude if hasattr(self.request, 'latitude') else ''
        # context['longitude'] = self.request.longitude if hasattr(self.request, 'longitude') else ''
        return context
    
@login_required
def export_eglises(request, format):
    eglises = Eglise.objects.select_related('ville', 'region', 'pasteur').all()
    print(eglises)
    
    if format == 'pdf':
        buffer = export_eglises_to_pdf(eglises)
        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="liste_eglises.pdf"'
    elif format == 'excel':
        buffer = export_eglises_to_excel(eglises)
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="liste_eglises.xlsx"'
    elif format == 'word':
        buffer = export_eglises_to_word(eglises)
        response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = 'attachment; filename="liste_eglises.docx"'
    else:
        return HttpResponse("Format non supporté", status=400)
    
    return response

class EgliseUpdateView(LoginRequiredMixin, UpdateView):
    model = Eglise
    fields = ['nom', 'pasteur', 'phone', 'image', 'region', 'ville', 'pays', 'latitude', 'longitude']   # Ou `fields = [...]` si pas de formulaire
    template_name = 'eglise/eglise_update.html'
    success_url = reverse_lazy('eglise-list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['latitude'] = f"{self.object.latitude:.2f}" if self.object.latitude is not None else None
        context['longitude'] = f"{self.object.longitude:.2f}" if self.object.longitude is not None else None
        return context
    
    def form_valid(self, form):

        messages.success(self.request, f"✅ {self.object.nom}' enregistrée avec succès.")
        return super().form_valid(form)
    


class EgliseDeleteView(LoginRequiredMixin,DeleteView):
    model = Eglise
    success_url = reverse_lazy('eglise-list')
    
    def delete(self, request, *args, **kwargs):
        self.object = self.get_object()
        self.object.delete()
        return JsonResponse({'success': True, 'redirect_url': str(self.success_url)})



# materiel_views.py

class MaterielListView(LoginRequiredMixin,ListView):
    model = Materiel
    template_name = 'materiel/materiel_list.html'
    context_object_name = 'materiel'
    paginate_by = 10

    def get_queryset(self, *args, **kwargs):
        user = self.request.user
        
        # Cas où l'utilisateur n'est pas connecté
        if not user.is_authenticated:
            messages.error(self.request, "Vous devez vous connecter pour voir vos materiels.")
            return Materiel.objects.none()
        
        # Cas sans église
        if not hasattr(user, 'eglise') or not user.eglise:
            messages.error(self.request, "Vous n’êtes pas rattaché à une église.")
            return Materiel.objects.none()

        # Cas sans logistique
        logistique = user.eglise.logistique_eglise
        # print("logistique", logistique)
        if not hasattr(user.eglise, 'logistique_eglise') or not user.eglise.logistique_eglise:
            messages.warning(self.request, "Aucune logistique associée à votre église. Faite une demande ici « Ajouter ».") 
            return Materiel.objects.none()

        # Par défaut on filtre par logistique de l'utilisateur et non supprimé
        qs = Materiel.objects.filter(logistique=logistique, is_deleted=False)

        # Messages informatifs
        if not qs.exists():
            messages.info(self.request, f'Il semble que vous n’ayez pas encore ajouté de matériel même si "{user.eglise.nom}" dispose d’une logistique.')
        else:
            messages.success(self.request, f'{user.eglise.nom} dispose d’une logistique.')

        # Si l'utilisateur est un responsable national
        if self.test_func():
            qs = Materiel.objects.filter(is_deleted=False)  # ne pas limiter à sa logistique

            # Filtres GET
            categorie = self.request.GET.get("categorie")
            sous_categorie = self.request.GET.get("sous_categorie")
            logistique_filter = self.request.GET.get("logistique")
            searchinput = self.request.GET.get("searchInput")

            if categorie:
                qs = qs.filter(categorie__nom=categorie)
            elif sous_categorie:
                qs = qs.filter(sous_categorie__nom=sous_categorie)
            elif logistique_filter:
                qs = qs.filter(logistique__id=logistique_filter)
            elif searchinput:
                qs = qs.filter(nom__icontains=searchinput)
            else:
                qs = Materiel.objects.filter(is_deleted=False)


        return qs.order_by('-id')

    def test_func(self):
        return self.request.user.groups.filter(name='Responsable Logistique National').exists()

    def get_context_data(self, **kwargs):
       
        context = super().get_context_data(**kwargs)
        context['nb_materiels'] = Materiel.objects.values_list('nom', flat=True).count()
        context['categories'] = CategorieMateriel.objects.all()
        context['sous_categories'] = SousCategorieMateriel.objects.all()
        context['logistiques'] = Logistique.objects.select_related('eglise')
        context['corbeille'] = Materiel.objects.filter(is_deleted=True)
        context['can_add_materiel'] = self.request.user.has_perm('logistque.add_materiel') or self.request.user.groups.filter(name='Responsable Logistique National').exists()
        if self.request.user.is_authenticated:
            context['derniere_demande'] = DemandePermission.objects.filter(
                user=self.request.user,
                permission_demande='logistque.add_materiel'
            ).order_by('-date_demande').first()

        return context
    
    def render_to_response(self, context, **response_kwargs):
        # Si requête AJAX, retourner les sous-catégories JSON
        if self.request.headers.get('x-requested-with') == 'XMLHttpRequest' and 'categorie' in self.request.GET:
            nom_categorie = self.request.GET.get('categorie')
            try:
                categorie = CategorieMateriel.objects.get(nom=nom_categorie)
                sous_cats = CategorieMateriel.objects.filter(categorie=categorie).values('id', 'nom')
                return JsonResponse(list(sous_cats), safe=False)
            except CategorieMateriel.DoesNotExist:
                return JsonResponse([], safe=False)
        return super().render_to_response(context, **response_kwargs)




class DemandePermissionCreateView(LoginRequiredMixin, CreateView):
    model = DemandePermission
    form_class = DemandePermissionForm
    template_name = 'materiel/demande_permission.html'
    success_url = reverse_lazy('materiel-list')

    def form_valid(self, form):
        # Refuser une autre demande en attente
        existing = DemandePermission.objects.filter(user=self.request.user, permission_demande='logistque.add_materiel', statut='en_attente')
        if existing.exists():
            messages.warning(self.request, "⏳ Vous avez déjà une demande en attente.")
            return redirect(self.success_url)

        form.instance.user = self.request.user
        form.instance.permission_demande = 'logistque.add_materiel'
    
        # Envoi d’e-mail
        try:
            # msg = MIMEText(f"Utilisateur : {self.request.user.get_full_name()} ({self.request.user.email})\n"
            #         f"Permission demandée : {form.instance.permission_demande}\n"
            #         f"Raison :\n{form.cleaned_data['raison']}")
            # msg["Subject"] = "📩 Nouvelle demande de permission"
            # msg["From"] = settings.EMAIL_HOST_USER
            # msg["To"] = self.request.user.email
            # with smtplib.SMTP("smtp.hostinger.com", 465) as server:
            #     server.starttls()
            #     server.login("logistique@sajholding.org", "Sajholding@2025")
            #     server.send_message(msg)
            pass
           
        except Exception as e:
            print("Erreur envoi mail :", e)  # utile pour le debug

        messages.success(self.request, "✅ Votre demande a été envoyée à l’administrateur.")
        
        return super().form_valid(form)
    

@csrf_exempt
@login_required
def export_materiels_excel(request):
    if request.method == "POST":
        try:
            body = json.loads(request.body)
            data = body.get("data", [])
        except Exception:
            return HttpResponse("Erreur de données", status=400)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Matériels filtrés"

        # Entêtes (doivent correspondre à ce qu'il y a dans le tableau HTML)
        headers = ['Nom', 'Catégorie', 'Sous-catégorie', 'Quantité', 'Logistique', 'Eglise']
        ws.append(headers)

        for row in data:
            ws.append(row)

        # Style simple
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=MaterielsFiltres.xlsx'
        wb.save(response)
        return response

    return HttpResponse("Méthode non autorisée", status=405)





class MaterielCreateView(LoginRequiredMixin, CreateView):
    model = Materiel
    fields = ['nom', 'categorie', 'sous_categorie', 'quantite', 'eglise', 'description']
    template_name = 'materiel/materiel_form_create.html'
    success_url = reverse_lazy('materiel-list')
    


    def dispatch(self, request, *args, **kwargs):
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return self.ajax_upload(request)

        if not request.user.is_authenticated:
            messages.error(request, "Vous devez vous connecter pour ajouter un matériel.")
            return redirect('login')

        if not hasattr(request.user, 'eglise') or not request.user.eglise:
            messages.error(request, "Vous n’êtes pas rattaché à une église.")
            return redirect('materiel-list')

        if request.user.role == 'membre':
            messages.error(request, "Vous devez être responsable ou pasteur pour ajouter un matériel.")
            return redirect('materiel-list')



        return super().dispatch(request, *args, **kwargs)

    def ajax_upload(self, request):
        if request.method != 'POST':
            return JsonResponse({'error': 'Méthode non autorisée'}, status=405)

        form = self.get_form()
        if form.is_valid():
            self.object = form.save(commit=False)
            self.object.logistique = request.user.eglise.logistique_eglise
            self.object.save()

            files = request.FILES.getlist('images')
            uploaded_files = []

            for file in files:
                img = MaterielImage.objects.create(
                    materiel=self.object,
                    image=file,
                    description=""
                )
                uploaded_files.append({
                    'id': img.id,
                    'url': img.image.url,
                    'name': img.image.name
                })

            return JsonResponse({'success': True, 'files': uploaded_files, 'materiel_id': self.object.id})
        return JsonResponse({'error': 'Formulaire invalide'}, status=400)

    def form_valid(self, form):
        form.instance._request = self.request
        form.instance.logistique = self.request.user.eglise.logistique_eglise
        
        if form.instance.eglise != self.request.user.eglise:
            messages.error(self.request, "Vous ne pouvez pas ajouter un matériel pour une autre église.")
            return JsonResponse({'error': 'Vous ne pouvez pas ajouter un matériel pour une autre église.'}, status=400)

        with transaction.atomic():
            self.object = form.save()

            if 'delete_images' in self.request.POST:
                delete_ids = [int(id) for id in self.request.POST.getlist('delete_images')]
                self.object.images_materiel.filter(id__in=delete_ids).delete()

            if 'images' in self.request.FILES:
                for image in self.request.FILES.getlist('images'):
                    MaterielImage.objects.create(
                        materiel=self.object,
                        image=image
                    )

        messages.success(self.request, "Matériel créé avec succès!")
        return super().form_valid(form)
    


    

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['is_create'] = True
        context['images'] = []
        return context

def get_sous_categories(request):
    categorie_id = request.GET.get('categorie')
    sous_categories = SousCategorieMateriel.objects.filter(categorie_id=categorie_id).values('id', 'nom')
    print(sous_categories)
    return JsonResponse(list(sous_categories), safe=False)
    
class MaterielUpdateView(LoginRequiredMixin,UpdateView):
    model = Materiel
    fields = ['nom', 'categorie', 'sous_categorie', 'quantite', 'eglise']
    template_name = 'materiel/materiel_form.html'
    success_url = reverse_lazy('materiel-list')

    def dispatch(self, request, *args, **kwargs):
        # Séparer le traitement des requêtes AJAX
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return self.ajax_upload(request)
        
        # Cas où l'utilisateur n'est pas connecté
        if not self.request.user.is_authenticated:
            messages.error(self.request, "Vous devez vous connecter pour modifier un materiel.")
            return Materiel.objects.none()

        if not hasattr(request.user, 'eglise') or not request.user.eglise:
            messages.error(request, "Vous n’ête pas rattaché à une église.")
            return redirect('materiel-list')
    
        if not User.objects.filter(role='responsable').exists():
            messages.error(request, "Votre église ne dispose pas encore d’une logistique.")
            return redirect('materiel-list')
        return super().dispatch(request, *args, **kwargs)

    def ajax_upload(self, request):
        if request.method != 'POST':
            return JsonResponse({'error': 'Méthode non autorisée'}, status=405)
            
        self.object = self.get_object()
        files = request.FILES.getlist('images')
        uploaded_files = []
        
        for file in files:
            img = MaterielImage.objects.create(
                materiel=self.object,
                image=file,
                description=""
            )
            uploaded_files.append({
                'id': img.id,
                'url': img.image.url,
                'name': img.image.name
            })
        
        return JsonResponse({'success': True, 'files': uploaded_files})

    def form_valid(self, form):
        form.instance._request = self.request
        with transaction.atomic():
            self.object = form.save()
            
            if 'delete_images' in self.request.POST:
                delete_ids = [int(id) for id in self.request.POST.getlist('delete_images')]
                self.object.images_materiel.filter(id__in=delete_ids).delete()
        
        messages.success(self.request, "Matériel mis à jour avec succès!")
        return super().form_valid(form)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['images'] = self.object.images_materiel.all()
        return context

class MaterielDeleteView(LoginRequiredMixin,DeleteView):
    model = Materiel
    success_url = reverse_lazy('materiel-list')

    def dispatch(self, request, *args, **kwargs):
        obj = self.get_object()
        print(obj.logistique.eglise.id)
        verify = Logistique.objects.filter(responsable=request.user.eglise.id).exists()
        
        if not verify or not obj.logistique.eglise.id == request.user.eglise.id:
            
            if request.headers.get('x-requested-with') == 'XMLHttpRequest':
                return JsonResponse({'error': "Suppression non autorisée"}, status=403)
            messages.error(request, "Vous ne pouvez pas supprimer ce matériel.")
            return redirect('materiel-list')
        return super().dispatch(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        # print("CSRF Token reçu :", request.META.get("HTTP_X_CSRFTOKEN"))
        self.object = self.get_object()
        self.object.is_deleted = True
        self.object.save()

        if request.headers.get('x-requested-with') == 'XMLHttpRequest':
            return JsonResponse({'success': True})
        messages.success(request, "Matériel déplacé dans la corbeille.")
        return redirect(self.success_url)

@login_required  
def ajax_delete(request, pk):
    if request.method == 'POST' and request.headers.get('x-requested-with') == 'XMLHttpRequest':
        materiel = get_object_or_404(Materiel, pk=pk)
        # Vérifie si l'utilisateur a le droit de supprimer
        # if not (hasattr(materiel.logistique, 'eglise') and materiel.logistique.eglise == request.user.eglise) and not request.user.is_superuser:
        #     return JsonResponse({'error': 'Ce matériel n\'appartient pas à votre église.'}, status=403)
        # Vérifie si le matériel appartient à une logistique d'une église
        verify = Logistique.objects.filter(eglise=materiel.logistique.eglise.id , responsable=request.user.eglise.id).exists()
        
        if verify:
            materiel.delete()
            return JsonResponse({'success': True})
        else:
            return JsonResponse({'error': 'Vous ne pouvez pas supprimer ce matériel.'}, status=403)
    return JsonResponse({'error': 'Méthode non autorisée'}, status=405)
class MaterielRestoreView(LoginRequiredMixin,View):
    def post(self, request, pk):
        materiel = get_object_or_404(Materiel, pk=pk)

        # Vérifie si l'utilisateur a le droit de restaurer
        # if not (hasattr(materiel.logistique, 'eglise') and materiel.logistique.eglise == request.user.eglise) and not request.user.is_superuser:
        #     return JsonResponse({'error': 'Ce matériel n\'appartient pas à votre église.'}, status=403)
        # Vérifie si le matériel appartient à une logistique d'une église
        verify = Logistique.objects.filter(eglise=materiel.logistique.eglise.id , responsable=request.user.eglise.id).exists()
        
        if verify:
            materiel.is_deleted = False
            materiel.save()
            return JsonResponse({'success': True, 'message': 'Matériel restauré avec succès'})
        return JsonResponse({'error': 'Ce matériel n\'appartient pas à la logistique de votre église pour être restauré.'}, status=403)

        


class MaterielDetailView(LoginRequiredMixin,DetailView):
    model = Materiel
    template_name = 'materiel/materiel_detail.html'
    
    def dispatch(self, request, *args, **kwargs):
        # Cas ou l'utilisateur n'est pas connecté
        if not request.user.is_authenticated:
            messages.error(request, "Vous devez vous connecter pour voir ce materiel.")
            return redirect('login')
        return super().dispatch(request, *args, **kwargs)


# evenement_views.py

class CampListView(LoginRequiredMixin, ListView):
    model = CampMondial
    template_name = 'camp/camp_list.html'
    context_object_name = 'camps'

    def get_queryset(self):
        qs = super().get_queryset()
        # Filtrer si besoin selon la région de l'utilisateur
        return qs

class CampDetailView(LoginRequiredMixin, DetailView):
    model = CampMondial
    template_name = 'camp/camp_detail.html'

class CampCreateView(LoginRequiredMixin, CreateView):
    model = CampMondial
    fields = ['titre', 'ville', 'date_debut', 'date_fin', 'description', 'renfort_national']
    template_name = 'camp/camp_form.html'
    success_url = reverse_lazy('camp-list')




class DashboardView(LoginRequiredMixin, TemplateView):
    template_name = "dashboard_client.html"
    
    def dispatch(self, request, *args, **kwargs):
        # Cas ou l'utilisateur n'est pas connecté
        if not request.user.is_authenticated:
            messages.error(request, "Connectez-vous pour acceder au Tableau de bord.")
            return redirect('login')
        return super().dispatch(request, *args, **kwargs)
    

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        today = timezone.now().date()

        # 1. Données de base communes
        self.add_base_context(context, user, today)
        
        # 2. Données pour les graphiques
        self.add_charts_data(context, user)
        
        # 3. Données spécifiques au rôle
        self.add_role_specific_context(context, user, today)

        return context

    def add_base_context(self, context, user, today):
        """Ajoute les données communes à tous les utilisateurs"""
        context.update({
            'today': today,
            'user': user,
            'derniere_demande': self.get_last_permission_request(user),
            'stats': self.get_global_stats()
        })

    def get_last_permission_request(self, user):
        """Récupère la dernière demande de permission"""
        if user.is_authenticated:
            return DemandePermission.objects.filter(
                user=user,
                permission_demande='logistque.add_materiel'
            ).order_by('-date_demande').first()
        return None

    def get_global_stats(self):
        """Récupère les statistiques globales"""
        return {
            'eglises_count': Eglise.objects.count(),
            'evenements_count': Evenement.objects.count(),
            'materiels_count': Materiel.objects.count(),
            'membres_count': MembreLogistique.objects.count(),
            'camp_count': CampMondial.objects.count(),
            'nb_regions': Region.objects.count(),
            'nb_villes': Ville.objects.count(),
            'nb_membres_forme': FormationLogisticien.objects.select_related('formations').filter(est_forme=True).count(),
            'nb_membres_non_forme': FormationLogisticien.objects.filter(est_forme=False).count(),
            'nb_responsables_logistique': User.objects.filter(role__in=['rln', 'rll']).count(),
            'nb_pasteurs': User.objects.filter(role__in=['pasteur_national', 'pasteur_local']).count(),
            'nb_techniciens': User.objects.filter(role='technicien').count()
        }

    def add_charts_data(self, context, user):
        """Prépare les données pour les graphiques"""
        context.update({
            'eglises_chart': self.get_eglises_chart_data(),
            'materiel_chart': self.get_materiel_chart_data(user),
            'events_chart': self.get_events_chart_data(),
            'map_data': self.get_map_data()
        })

    def get_eglises_chart_data(self):
        """Données pour le graphique des églises par région"""
        regions_data = Region.objects.annotate(
            nb_eglises=Count('region_eglise', distinct=True)
        ).order_by('nom')
        
        return {
            'labels': [r.nom for r in regions_data],
            'data': [r.nb_eglises for r in regions_data],
            'colors': self.generate_colors(len(regions_data))
        }

    def get_materiel_chart_data(self, user):
        """Données pour le graphique du matériel"""
        materiel_query = Materiel.objects.all()
        if hasattr(user, 'responsable_logistique'):
            materiel_query = materiel_query.filter(logistique=user.responsable_logistique)
        
        materiel_data = materiel_query.values(
            'categorie__nom'
        ).annotate(
            total=Sum('quantite')
        ).order_by('categorie__nom')

        return {
            'labels': [m['categorie__nom'] for m in materiel_data],
            'data': [m['total'] for m in materiel_data],
            'colors': self.generate_colors(len(materiel_data), palette='material')
        }

    def get_events_chart_data(self):
        """Données pour le graphique des événements"""
        events_data = Evenement.objects.values(
            'type_evenement'
        ).annotate(
            count=Count('id')
        ).order_by('type_evenement')

        return {
            'labels': [dict(Evenement.TYPE_EVENEMENT)[e['type_evenement']] for e in events_data],
            'data': [e['count'] for e in events_data],
            'colors': self.generate_colors(len(events_data), palette='events')
        }

    def add_role_specific_context(self, context, user, today):
        """Ajoute les données spécifiques au rôle de l'utilisateur"""
        if user.role in ['rln', 'super_admin', 'pasteur_national']:
            self.add_responsable_national_context(context, user, today)
        elif user.role in ['rll', 'pasteur_local']:
            self.add_responsable_local_context(context, user, today)
        elif user.role == 'technicien':
            self.add_technicien_context(context, user, today)

    def get_map_data(self):
        """Données pour le rendu de la carte de Côte d'Ivoire"""
        eglises_avec_loc = Eglise.objects.exclude(latitude__isnull=True).exclude(longitude__isnull=True)
        map_points = []
        for eglise in eglises_avec_loc:
            dispo = Materiel.objects.filter(eglise=eglise, etat='OP').aggregate(Sum('quantite'))['quantite__sum'] or 0
            map_points.append({
                'nom': eglise.nom,
                'lat': eglise.latitude,
                'lon': eglise.longitude,
                'disponibilite': dispo,
                'ville': eglise.ville.nom if eglise.ville else ""
            })
        return map_points

    def add_technicien_context(self, context, user, today):
        """Contexte spécifique pour les techniciens"""
        # Un technicien voit son planning et les pannes qu'il a signalées
        context.update({
            'role': 'Membre Technicien',
            'mes_taches': ChronogrammeItem.objects.filter(responsable=user, evenement__date_fin__gte=today).order_by('heure_debut'),
            'mes_pannes': FicheDefectuosite.objects.filter(rapporteur=user).order_by('-date_signalement')[:5]
        })

    def add_responsable_local_context(self, context, user, today):
        """Contexte spécifique pour les responsables locaux"""
        logistique = user.responsable_logistique
        ville = logistique.eglise.ville
        materiels = Materiel.objects.filter(logistique=logistique)
        eglises = Eglise.objects.filter(ville=ville)
        logistiques = Logistique.objects.filter(eglise__in=eglises)

        context.update({
            'role': 'Responsable Logistique Local',
            'ville_responsable': ville.nom,
            'local_stats': {
                'materiel_total': materiels.aggregate(Sum('quantite'))['quantite__sum'] or 0,
                'nb_eglises': eglises.count(),
                'nb_membres': MembreLogistique.objects.filter(logistique__in=logistiques).count(),
                'nb_regions': Region.objects.count(),
                'nb_villes': Ville.objects.count(),
                'materiels_disponibles': materiels.filter(quantite__gt=0).count(),
                'materiels_manquants': materiels.filter(quantite=0).count(),
                'materiels_en_panne': materiels.filter(etat='PA').count(),
                #'camps_count': CampMondial.objects.filter(ville=ville, date_debut__gte=today).all().count(),
                'renforts_count': CampMondial.objects.filter(renfort_national=True, date_debut__gte=today).all().count(),
                'camps_count': CampMondial.objects.filter(date_debut__gte=today).count(),
                'evenements_a_venir': CampMondial.objects.filter(
                ville=ville,
                date_debut__gte=today
            ).order_by('date_debut')[:5],
                
                
            },
            
            'materiel_details': self.get_materiel_details(logistique),
            'renforts': CampMondial.objects.filter(renfort_national=True, date_debut__gte=today).all().order_by('-date_debut'),
            'prochains_evenements': Evenement.objects.filter(
                date_debut__gte=today,
            ).order_by('date_debut')[:5]
        })

    def get_materiel_details(self, logistique):
        """Détails du matériel pour les responsables locaux"""
        return {
            'par_categorie': Materiel.objects.filter(logistique=logistique)
                .values('categorie__nom')
                .annotate(total=Sum('quantite')),
            'par_sous_categorie': Materiel.objects.filter(logistique=logistique)
                .values('sous_categorie__nom')
                .annotate(total=Sum('quantite'))
        }

    def add_responsable_national_context(self, context, user, today):
        """Contexte spécifique pour les responsables nationaux"""
        context.update({
            'role': 'Responsable Logistique National',
            'national_stats': {
                'camps_count': CampMondial.objects.filter(date_debut__gte=today).count(),
                'renforts_count': CampMondial.objects.filter(
                    renfort_national=True, 
                    date_debut__gte=today
                ).count(),
                'materiel_national': Materiel.objects.filter(
                    logistique__eglise__nom="Eglise Nationale"
                ).aggregate(Sum('quantite'))['quantite__sum'] or 0
            },
            'camps': CampMondial.objects.filter(date_debut__gte=today),
            'renforts_demandes': CampMondial.objects.filter(
                renfort_national=True, 
                date_debut__gte=today
            ),
        })

    def add_pasteur_context(self, context, user, today):
        """Contexte spécifique pour les pasteurs"""
        context.update({
            'role': 'Pasteur de la section',
            'pasteur_stats': {
                'nb_eglises': Eglise.objects.count(),
                'camps_count': CampMondial.objects.filter(date_debut__gte=today).count()
            },
            'prochains_evenements': Evenement.objects.filter(
                date_debut__gte=today
            ).all().order_by('-date_debut')[:5]
        })
        print(context['pasteur_stats'].keys)

    def generate_colors(self, count, palette='default'):
        """Génère des palettes de couleurs cohérentes"""
        palettes = {
            'default': ['#FF6384', '#36A2EB', '#FFCE56', '#4BC0C0', '#9966FF'],
            'material': ['#FF5252', '#FF4081', '#E040FB', '#7C4DFF', '#536DFE'],
            'events': ['#FF9F40', '#FFCD56', '#4BC0C0', '#9966FF', '#FF6384']
        }
        return palettes[palette][:count]
    
    
class MaterielStatsAPIView(LoginRequiredMixin,View):
    def get(self, request):
        group_by = request.GET.get('group_by', 'categorie')
        user = request.user
        
        # Base queryset avec filtrage par utilisateur si nécessaire
        materiel_query = Materiel.objects.all()
        if hasattr(user, 'responsable_logistique'):
            materiel_query = materiel_query.filter(logistique=user.responsable_logistique)
        
        # Groupement des données
        if group_by == 'sous_categorie':
            data = self.get_sous_categorie_stats(materiel_query)
        else:
            data = self.get_categorie_stats(materiel_query)
        
        return JsonResponse(data)

    def get_categorie_stats(self, queryset):
        stats = (
            queryset
            .values('categorie__nom')
            .annotate(
                total=Sum('quantite'),
                count=Count('id')
            )
            .order_by('categorie__nom')
        )
        
        return {
            'labels': [item['categorie__nom'] for item in stats],
            'values': [item['total'] for item in stats],
            'colors': self.generate_colors(len(stats), palette='material')
        }

    def get_sous_categorie_stats(self, queryset):
        stats = (
            queryset
            .values('sous_categorie__nom')
            .annotate(
                total=Sum('quantite'),
                count=Count('id')
            )
            .order_by('sous_categorie__nom')
        )
        
        return {
            'labels': [item['sous_categorie__nom'] for item in stats],
            'values': [item['total'] for item in stats],
            'colors': self.generate_colors(len(stats), palette='material')
        }

    def generate_colors(self, count, palette='material'):
        palettes = {
            'material': ['#FF5252', '#FF4081', '#E040FB', '#7C4DFF', '#536DFE'],
            'default': ['#FF6384', '#36A2EB', '#FFCE56', '#4BC0C0', '#9966FF']
        }
        return palettes[palette][:count]
    
def login_view(request):
    phone = request.session.pop('signup_phone', None)
    if request.method == 'POST':
        identifiant = request.POST.get('user')
        password = request.POST.get('password')
        remember = request.POST.get('remember_me')
        
        

        # Recherche utilisateur par email ou téléphone
        try:
            user = User.objects.get(email=identifiant)
        except User.DoesNotExist:
            try:
                user = User.objects.get(phone=identifiant)
            except User.DoesNotExist:
                user = None

        if user:
            user = authenticate(request, username=user, password=password)
            if user:
                login(request, user)
                if not remember:
                    request.session.set_expiry(0)  # expire à la fermeture
                return redirect('dashboard-client')
        
        messages.error(request, "Identifiants incorrects")
    return render(request, 'authenticate/login.html', {'phone': phone})
    
    
def logout_view(request):
    logout(request)
    return redirect('login')

@login_required
def logout_confirm(request):
    if request.method == 'POST':
        logout(request)
        return redirect('login')  # nom de ta page de connexion
    return render(request, 'registration/confirm_logout.html')

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

class PackingListPDFView(LoginRequiredMixin, View):
    def get(self, request, event_id):
        event = get_object_or_404(Evenement, pk=event_id)
        response = HttpResponse(content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="liste_colisage_{event_id}.pdf"'

        doc = SimpleDocTemplate(response, pagesize=A4)
        elements = []
        styles = getSampleStyleSheet()

        elements.append(Paragraph(f"Liste de Colisage : {event.titre}", styles['Title']))
        elements.append(Paragraph(f"Date : {event.date_debut.strftime('%d/%m/%Y') if event.date_debut else ''}", styles['Normal']))
        elements.append(Paragraph(f"Lieu : {event.lieu if event.lieu else 'Non spécifié'}", styles['Normal']))
        elements.append(Paragraph("<br/><br/>", styles['Normal']))

        data = [["Matériel", "Quantité", "Catégorie", "État"]]
        for em in event.materiels_associes.all():
            data.append([
                em.materiel.nom if em.materiel else "Inconnu",
                str(em.quantite),
                em.materiel.categorie.nom if em.materiel and em.materiel.categorie else "",
                em.materiel.get_etat_display() if em.materiel else ""
            ])

        if len(data) == 1:
            data.append(["Aucun matériel", "-", "-", "-"])

        t = Table(data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(t)
        doc.build(elements)
        return response
